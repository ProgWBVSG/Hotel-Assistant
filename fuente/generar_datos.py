# -*- coding: utf-8 -*-
"""Convierte el Excel real en un archivo JS que la app carga directamente."""
import openpyxl, re, datetime, json, collections

wb = openpyxl.load_workbook('2026.xlsx', data_only=True)
AREAS = {'ird':'In Room Dining','in room dining':'In Room Dining',
         'exchange lane':'Exchange Lane','penny blue':'Penny Blue'}

def mapa(ws):
    m = {}
    for rng in ws.merged_cells.ranges:
        v = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row+1):
            for c in range(rng.min_col, rng.max_col+1):
                m[(r,c)] = v
    return m

def anchos(ws):
    """Filas combinadas que ocupan casi todo el ancho -> comentarios."""
    out = {}
    for rng in ws.merged_cells.ranges:
        if rng.max_col - rng.min_col >= 4:
            out[rng.min_row] = ws.cell(rng.min_row, rng.min_col).value
    return out

def fecha_de(nombre, ws, m):
    s = re.sub(r'\s*\(\d\)', '', nombre).replace('..','.').strip()
    mm = re.match(r'^(\d{1,2})\.(\d{1,2})\.(\d{2,4})$', s)
    if mm:
        d, mo, y = int(mm.group(1)), int(mm.group(2)), int(mm.group(3))
        if y < 100: y += 2000
        try: return datetime.date(y, mo, d)
        except: pass
    for r in range(1, 12):
        if str(m.get((r,1), ws.cell(r,1).value) or '').strip().lower() == 'date':
            v = m.get((r,2), ws.cell(r,2).value)
            if isinstance(v, datetime.datetime): return v.date()
            if isinstance(v, datetime.date): return v
            if isinstance(v, (int,float)):
                s2 = str(int(v)).zfill(8)
                try: return datetime.date(int(s2[4:]), int(s2[2:4]), int(s2[:2]))
                except: pass
    return None


PAT_HORA = re.compile(r'^\s*(\d{1,2})[:.h](\d{2})\s*[-–—]\s*(\d{1,2})[:.h](\d{2})\s*$')

def parsear_turno(nombre, horario, nota):
    if not isinstance(horario, str): return None
    m = PAT_HORA.match(horario.replace(' ', '').replace('-', ' - ').replace('  ', ' '))
    if not m:
        m = PAT_HORA.match(re.sub(r'\s+', '', horario).replace('-', ' - '))
    if not m: return None
    h1, m1, h2, m2 = int(m.group(1)), int(m.group(2)), int(m.group(3)), int(m.group(4))
    if h1 > 23 or h2 > 23 or m1 > 59 or m2 > 59: return None
    ini = h1 * 60 + m1
    fin = h2 * 60 + m2
    if fin <= ini: fin += 24 * 60          # turno que cruza la medianoche
    dur = fin - ini
    if dur <= 0 or dur > 16 * 60: return None
    descanso = 0
    if isinstance(nota, str):
        mb = re.search(r'(\d{2,3})\s*m', nota.lower())
        if mb and 'no break' not in nota.lower():
            descanso = int(mb.group(1))
    return {'quien': str(nombre).strip()[:28], 'desde': ini, 'hasta': fin,
            'descanso': descanso, 'horas': round((dur - descanso) / 60.0, 2)}

dias = {}
for nombre in wb.sheetnames:
    ws = wb[nombre]; m = mapa(ws); com_filas = anchos(ws)
    fecha = fecha_de(nombre, ws, m)
    if not fecha: continue
    clave = str(fecha)

    area = None; serv = None
    lineas = collections.defaultdict(dict)
    comentarios = []
    orden_areas = []
    turnos = []

    for r in range(1, ws.max_row+1):
        a = m.get((r,1), ws.cell(r,1).value)
        b = m.get((r,2), ws.cell(r,2).value)
        c = m.get((r,3), ws.cell(r,3).value)
        d = ws.cell(r,4).value
        a_s = str(a).strip() if a else ''
        b_s = str(b).strip() if b else ''
        c_s = str(c).strip() if c else ''

        if r in com_filas:
            t = com_filas[r]
            if isinstance(t, str) and len(t.strip()) > 25:
                comentarios.append({'texto': t.strip(), 'area': area, 'fila': r})
            continue

        if a_s.lower() in AREAS:
            area = AREAS[a_s.lower()]
            if area not in orden_areas: orden_areas.append(area)
        if b_s and b_s.lower() not in ('breakfast','lunch','dinner','overnight','all day') and not area:
            pass
        if b_s: serv = b_s
        if area and serv and c_s and isinstance(d, (int, float)):
            lineas[(area, serv)][c_s] = round(float(d), 2)

        # columnas E/F/G: turnos del personal
        e_v = ws.cell(r,5).value; f_v = ws.cell(r,6).value; g_v = ws.cell(r,7).value
        if isinstance(e_v, str) and e_v.strip() and len(e_v.strip()) < 30:
            t = parsear_turno(e_v, f_v, g_v)
            if t:
                t['area'] = area
                turnos.append(t)

    if not orden_areas: continue

    areas_out = {}
    for (ar, sv), mets in lineas.items():
        areas_out.setdefault(ar, {})[sv] = mets

    # dia ya existente: nos quedamos con el que tenga mas datos
    if clave in dias and len(json.dumps(dias[clave])) > len(json.dumps(areas_out)):
        continue

    dias[clave] = {
        'fecha': clave, 'hoja': nombre, 'areas': areas_out,
        'comentarios': comentarios, 'turnos': turnos
    }

salida = [dias[k] for k in sorted(dias)]
js = 'var DATOS_REALES = ' + json.dumps(salida, ensure_ascii=False, indent=0, separators=(',',':')) + ';\n'
open('../app/js/datos-reales.js','w',encoding='utf-8').write(js)

print('DIAS:', len(salida))
print('COMENTARIOS:', sum(len(d['comentarios']) for d in salida))
print('TURNOS:', sum(len(d['turnos']) for d in salida))
meses = collections.Counter(d['fecha'][:7] for d in salida)
print('MESES:', ', '.join('%s=%d' % (k, v) for k, v in sorted(meses.items())))
print('PESO:', round(len(js)/1024), 'KB')
